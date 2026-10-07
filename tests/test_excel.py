import numpy as np
import pytest
from openpyxl import load_workbook

from fish_viewer.models import Surface
from fish_viewer.workflow.volumes import surface_volume
from fish_viewer.workflow.excel import write_excel


def box_surface():
    vertices = np.array([[0,0,0], [2,0,0], [2,3,0], [0,3,0],
                         [0,0,4], [2,0,4], [2,3,4], [0,3,4]], dtype=float)
    faces = np.array([[0,2,1], [0,3,2], [4,5,6], [4,6,7], [0,1,5], [0,5,4],
                      [1,2,6], [1,6,5], [2,3,7], [2,7,6], [3,0,4], [3,4,7]])
    return Surface('Predicted cell 1', vertices, faces)


def test_mesh_volume_is_physical_translation_invariant_and_rejects_open_mesh():
    surface = box_surface()
    assert surface_volume(surface) == pytest.approx(24)
    surface.vertices += [100, -70, 20]
    surface.triangles[::2] = surface.triangles[::2, ::-1]  # Inconsistent input winding is repaired.
    assert surface_volume(surface) == pytest.approx(24)
    surface.triangles = surface.triangles[:-1]
    assert surface_volume(surface) is None


def test_workbook_keeps_numeric_distances_volumes_channel_counts_and_safe_text(tmp_path):
    report = dict(image='scan.ims', channels=[dict(name='Membrane'),dict(name='=BAD()'),dict(name='Red')],
        points=[dict(signal_id=1,cell_id=4,channel=1,score=.999,xyz=[1,2,3],border_xyz=[1,2,4],distance_um=1),
                dict(signal_id=2,cell_id=4,channel=2,score=.998,xyz=[2,3,4],border_xyz=[2,3,6],distance_um=2)],
        counts=[dict(cell_id=4,channel=1,channel_name='=BAD()',count=1),dict(cell_id=4,channel=2,channel_name='Red',count=1)],
        cell_volumes=[dict(cell_id=4,volume_um3=24,voxel_volume_um3=25,volume_status='closed_mesh')])
    path = write_excel(tmp_path/'report.xlsx', report)
    workbook = load_workbook(path)
    assert workbook['Cells'].max_row == 2
    assert workbook['Cells']['B2'].value == 24 and workbook['Cells']['D2'].value == 2
    assert workbook['Signals']['M2'].value == 1 and workbook['Signals']['M3'].value == 2
    assert workbook['Signals']['C2'].value == 2 and workbook['Signals']['C3'].value == 3
    assert workbook['Signals']['D2'].data_type == 's'
    assert workbook['Channel counts'].max_row == 3
    workbook.close()


def test_degenerate_imaris_pole_triangles_do_not_bias_volume():
    surface = box_surface()
    surface.vertices = np.vstack([surface.vertices, surface.vertices[[0,0,1]]])
    surface.triangles = np.vstack([surface.triangles, [8,9,10]])
    assert surface_volume(surface) == pytest.approx(24)
