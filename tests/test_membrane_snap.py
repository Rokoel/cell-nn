from dataclasses import replace
import json
from types import SimpleNamespace

import numpy as np
import pytest
from scipy import ndimage as ndi

from fish_viewer.refinement.membrane import MembraneConfig, apply, snap, snap_cell


def sphere(spacing=(.15, .15, .15), shape=(56, 56, 56), radius=2.55, rim=2.72):
    grid = np.indices(shape)
    center = (np.asarray(shape) - 1) / 2
    distance = np.linalg.norm((grid - center[:, None, None, None]) * np.asarray(spacing)[:, None, None, None], axis=0)
    labels = (distance < radius).astype(np.uint32)
    image = np.exp(-((distance - rim) / .1) ** 2).astype(np.float32)
    return labels, SimpleNamespace(image=image, spacing=spacing), distance < rim


def test_local_peak_recovers_an_undersized_cell_and_bounds_motion():
    labels, prepared, truth = sphere()
    result, diagnostics = snap(labels, prepared)
    d = diagnostics[0]
    assert d['accepted'] and result.sum() > labels.sum()
    assert np.count_nonzero((result > 0) != truth) < np.count_nonzero((labels > 0) != truth)
    assert 0 < d['max_vertex_shift_um'] <= MembraneConfig().max_shift_um * MembraneConfig().strength
    assert set(np.unique(result)) == {0, 1}


def test_normal_peak_search_also_corrects_oversized_cells():
    labels, prepared, truth = sphere(radius=2.72, rim=2.55)
    result, diagnostic = snap(labels, prepared)
    assert diagnostic[0]['accepted'] and diagnostic[0]['mean_signed_shift_um'] < 0
    assert 0 < result.sum() < labels.sum()
    assert np.count_nonzero((result > 0) != truth) < np.count_nonzero((labels > 0) != truth)


@pytest.mark.parametrize('brightness', [0., .5, 1.])
def test_blank_or_flat_cell_channel_does_not_expand_contours(brightness):
    labels, prepared, _ = sphere()
    prepared.image.fill(brightness)
    result, diagnostics = snap(labels, prepared)
    assert np.array_equal(labels, result)
    assert not diagnostics[0]['accepted']


def test_refinement_preserves_neighbor_voxels_and_instance_identity():
    labels, prepared, _ = sphere()
    labels[20:36, 20:36, 45:49] = 7
    # A neighbor just outside the too-small mask blocks the proposed expansion.
    result, diagnostics = snap(labels, prepared)
    assert set(np.unique(result)) == {0, 1, 7}
    assert not np.any(result[labels == 7] == 1)
    assert not diagnostics[0]['accepted']
    assert 'overlap' in diagnostics[0]['fallback_reason']


def test_anisotropic_voxel_calibration_is_respected():
    labels, prepared, truth = sphere(spacing=(.3, .12, .12), shape=(36, 64, 64))
    result, diagnostics = snap(labels, prepared)
    assert diagnostics[0]['accepted']
    assert np.count_nonzero((result > 0) != truth) < np.count_nonzero((labels > 0) != truth)


def test_nonconvex_shape_is_preserved_instead_of_becoming_a_sphere():
    labels, prepared, _ = sphere(radius=2.8, rim=2.8)
    grid = np.indices(labels.shape)
    indentation = np.linalg.norm((grid - np.array([27.5, 27.5, 46])[:, None, None, None]) * .15, axis=0) < 1.4
    truth = (labels > 0) & ~indentation
    # An eroded contour should move toward the visible indented surface locally.
    original = ndi.binary_erosion(truth).astype(np.uint32)
    boundary = truth & ~ndi.binary_erosion(truth)
    prepared.image = np.exp(-(ndi.distance_transform_edt(~boundary, sampling=.15) / .1) ** 2).astype(np.float32)
    result, diagnostics = snap(original, prepared, MembraneConfig(strength=1.))
    assert diagnostics[0]['accepted']
    assert not result[27, 27, 46]  # deep indentation remains
    assert np.count_nonzero((result > 0) != truth) < np.count_nonzero((original > 0) != truth)


def test_invalid_settings_and_mismatched_grid_are_rejected():
    with pytest.raises(ValueError, match='finite'):
        replace(MembraneConfig(), max_shift_um=float('nan')).validate()
    with pytest.raises(ValueError, match='shift'):
        replace(MembraneConfig(), max_shift_um=2).validate()
    with pytest.raises(ValueError, match='grid'):
        snap(np.zeros((2, 2, 2), np.uint32), SimpleNamespace(image=np.zeros((3, 3, 3))))


def test_saved_run_preserves_baseline_geometry_and_refuses_cumulative_snapping(tmp_path):
    labels, prepared, _ = sphere()
    original = labels.copy()
    np.save(tmp_path / 'image.npy', prepared.image)
    metadata = dict(version=1, image='scan.ims', time=0, channel=1, channel_scores=[],
                    fingerprint='same', spacing_zyx=list(prepared.spacing), origin_xyz=[10, 20, 30],
                    cells=[dict(label=1, detection_score=.9)], refinement=dict(model='unchanged'))
    path = tmp_path / 'contours.npz'
    np.savez_compressed(path, labels=labels, metadata=np.array(json.dumps(metadata)))
    source = tmp_path / 'report.json'
    source.write_text(json.dumps(dict(status='complete', files=[dict(image='scan.ims', predictions=str(path))])))
    audit = tmp_path / 'audit.json'
    audit.write_text(json.dumps(dict(status='complete', cases=[dict(image='scan.ims', directory=str(tmp_path),
        audit=dict(fingerprint='same', channel=1, spacing_zyx=list(prepared.spacing), origin_xyz=[10, 20, 30]))])))
    report = apply(source, tmp_path / 'runs', audit=audit)
    assert report['status'] == 'complete', report.get('error')
    with np.load(path) as f:
        assert np.array_equal(f['labels'], original)
    with np.load(report['files'][0]['predictions']) as f:
        m = json.loads(str(f['metadata']))
        assert m['origin_xyz'] == [10, 20, 30] and m['channel'] == 1
        assert m['cells'][0]['detection_score'] == .9 and m['refinement']['model'] == 'unchanged'
        assert np.isclose(m['cells'][0]['volume_um3'], np.count_nonzero(f['labels']) * np.prod(prepared.spacing))
    duplicate = apply(report['directory'] + '/report.json', tmp_path / 'runs', audit=audit)
    assert duplicate['status'] == 'failed' and 'already been snapped' in duplicate['error']
    broken = json.loads(audit.read_text())
    broken['cases'][0]['audit']['origin_xyz'][0] += 1
    audit.write_text(json.dumps(broken))
    failed = apply(source, tmp_path / 'runs', audit=audit)
    assert failed['status'] == 'failed' and 'different physical grid' in failed['error']


def test_excel_remains_two_lists_with_membrane_snap_metadata(tmp_path):
    from openpyxl import load_workbook
    from fish_viewer.workflow.excel import write_excel
    report = dict(points=[], counts=[], cell_volumes=[dict(cell_id=1, volume_um3=10.)],
                  membrane_snap=dict(config=dict(max_shift_um=.3), cells=[dict(cell_id=1, accepted=True, mean_shift_um=.05)]))
    book = load_workbook(write_excel(tmp_path / 'report.xlsx', report), read_only=True, data_only=True)
    try:
        assert book.sheetnames == ['Cell volumes', 'Signal distances']
        assert list(book['Cell volumes'].values)[1] == ('NN',1,10.)
    finally:
        book.close()


@pytest.mark.gui
def test_analysis_inserts_optional_membrane_step_before_signals(qapp, tmp_path, monkeypatch):
    import fish_viewer.analysis.controller as module
    from fish_viewer.analysis.controller import AnalysisController
    monkeypatch.setattr(module, 'check_models', lambda *a: dict(engine='cell-center-box3d'))
    import fish_viewer.refinement.config as boundary
    monkeypatch.setattr(boundary, 'read_model', lambda *a: {})
    image = tmp_path / 'scan.ims'
    image.write_bytes(b'sample')
    controller = AnalysisController()
    monkeypatch.setattr(controller, '_launch', lambda *a: None)
    controller.start(image, tmp_path / 'runs', tmp_path / 'cells', tmp_path / 'signals',
                     refinement=tmp_path / 'crop', membrane_snap=dict(max_shift_um=.2))
    assert [s[0] for s in controller.stages] == ['cells', 'snap', 'signals', 'filter', 'measure', 'export']
    assert controller.state['membrane_snap']['max_shift_um'] == .2
