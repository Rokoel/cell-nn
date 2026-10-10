"""A single measurement sheet containing cell volumes and signal distances."""
from pathlib import Path
import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def write_excel(path, report):
    from .manual import manual_measurements
    manual = manual_measurements(report)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, headers, rows):
        page = workbook.create_sheet(name)
        page.append(headers)
        for row in rows:
            page.append(row)
            for cell in page[page.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                elif isinstance(cell.value, float):
                    cell.number_format = '0.0000'
        page.freeze_panes = 'A2'
        page.auto_filter.ref = page.dimensions
        for cell in page[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='205E58')
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = min(55, max(16, max(len(str(c.value or '')) for c in column) + 2))

    rows=[]
    for source, cells, points in [('Manual', manual['cells'], manual['points']),
                                  ('Program', report.get('cell_volumes', []), report.get('points', []))]:
        volumes={c['cell_id']:c.get('volume_um3') for c in cells}
        covered=set()
        for point in points:
            channel=point.get('channel')
            name=point.get('channel_name')
            if name is None and channel is not None:
                name=report['channels'][channel]['name']
            cell=point.get('cell_id'); covered.add(cell)
            rows.append([source,cell,volumes.get(cell),point['signal_id'],
                         channel+1 if channel is not None else None,name,point.get('distance_um')])
        rows.extend([source,c['cell_id'],c.get('volume_um3'),None,None,None,None] for c in cells if c['cell_id'] not in covered)
    sheet('Measurements', ['Source','Cell','Volume (µm³)','Signal','Channel (1-based)','Channel name','Distance (µm)'], rows)
    fd, temporary = tempfile.mkstemp(suffix='.xlsx', prefix='.fish-', dir=path.parent)
    os.close(fd)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        Path(temporary).unlink(missing_ok=True)
        workbook.close()
    return path
