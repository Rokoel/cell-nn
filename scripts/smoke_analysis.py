"""Exercise the focused app with actual models, export, A/B geometry and 2D/3D."""
import argparse
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import sys
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
os.environ['QT_API'] = 'pyside6'
if sys.platform.startswith('linux'):
    os.environ['QT_QPA_PLATFORM'] = 'xcb'

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--image', default='artifacts/imaris-compatibility/other-scans/170612_HAG50-2_Lam633m_K9546a_unannotated.ims')
parser.add_argument('--device', default='auto', choices=['auto', 'cpu', 'gpu'])
parser.add_argument('--package-root', type=Path, help='Verify the distributed source/model bundle')
parser.add_argument('--boundary-model', type=Path, help='Exercise the optional seeded 3D U-Net stage')
parser.add_argument('--cell-model', type=Path, help='Select StarDist or the center-and-box detector')
parser.add_argument('--signal-model', type=Path, help='Exercise a selected signal checkpoint')
parser.add_argument('--membrane-snap', action='store_true', help='Exercise local membrane peak refinement')
parser.add_argument('--cell-recovery', action='store_true', help='Exercise local membrane, retry and signal-guided cell recovery')
args = parser.parse_args()
if args.package_root:
    sys.path.insert(0, str(args.package_root.resolve()))
import numpy as np
from PySide6 import QtCore, QtWidgets
from fish_viewer.analysis.app import AnalysisWindow
from fish_viewer.workflow.distances import SurfaceLocator
output = Path('artifacts/analysis-smoke') / datetime.now().strftime('%Y%m%d-%H%M%S')
output.mkdir(parents=True)
app = QtWidgets.QApplication([])
app.setStyle('Fusion')
settings = QtCore.QSettings(str((output / 'settings.ini').resolve()), QtCore.QSettings.Format.IniFormat)
window = AnalysisWindow(settings=settings)
window.resize(1100, 720)
window.show()
errors = []
window.error = errors.append
last_progress = None


def progress(event):
    global last_progress
    text = f'{event["stage"]}/{event.get("stages", 5)} {event["phase"]} {event.get("current", 0)}/{event.get("total", 0)}'
    if text != last_progress:
        print(text, flush=True)
        last_progress = text


window.controller.progress.connect(progress)


def wait(predicate, seconds=900):
    deadline = time.monotonic()+seconds
    while not predicate() and time.monotonic()<deadline:
        app.processEvents()
        assert not errors, errors
        time.sleep(.02)
    assert predicate(), 'Timed out'
    assert not errors, errors


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def exercise():
    try:
        original_hash = digest(args.image)
        window.output.setText(str((output / 'runs').resolve()))
        window.device.setCurrentText(args.device)
        window.membrane_snap.setChecked(args.membrane_snap)
        window.cell_recovery.setChecked(args.cell_recovery)
        if args.membrane_snap:
            window.options_button.setChecked(True)
        if args.boundary_model:
            window.boundary_model.setText(str(args.boundary_model.resolve()))
        if args.cell_model:
            window.cell_model.setText(str(args.cell_model.resolve()))
        if args.signal_model:
            window.signal_model.setText(str(args.signal_model.resolve()))
        if args.package_root:
            assert Path(window.cell_model.text()).is_relative_to(args.package_root.resolve())
            assert Path(window.signal_model.text()).is_relative_to(args.package_root.resolve())
        window.add_paths([args.image])
        wait(lambda: window.dataset is not None and not window.loading)
        window.grab().save(str(output / 'ready.png'))
        window.run_button.click()
        assert window.running
        wait(lambda: not window.running and not window.loading)
        assert window.result_path
        report = window.result_report
        if args.cell_recovery:
            stages=json.loads(window.result_path.read_text())['stages']
            assert 'guide_signals' in stages and 'recover' in stages
            assert report['cell_recovery']['signal_guided_summary']
        if args.membrane_snap:
            assert report['membrane_snap']['cells'], 'Missing membrane peak diagnostics'
            assert 'snap' in json.loads(window.result_path.read_text())['stages']
        if window.boundary_model.text():
            assert report['boundary_refinement']['cells'], 'Missing boundary diagnostics'
            stages=json.loads(window.result_path.read_text())['stages']
            assert 'refine' in stages or report.get('cell_detector',{}).get('engine')=='cell-center-box3d'
        assert report['points'], 'Smoke scan must exercise real A/B measurements'
        assert window.distances.rowCount() == len(report['points'])
        assert sum(row['count'] for row in report['counts']) == len(report['points'])
        assert window.excel_button.isEnabled()
        assert len(report['cell_volumes']) == len(report['signal_filter']['kept_cell_ids'])
        assert all(row['volume_um3'] > 0 for row in report['cell_volumes'])
        from openpyxl import load_workbook
        exported = json.loads(window.result_path.read_text())
        workbook_path = next(window.result_path.parent / p for p in exported['files'] if p.endswith('.xlsx'))
        workbook = load_workbook(workbook_path, data_only=True)
        assert workbook.sheetnames == ['Cell volumes', 'Signal distances']
        nn_distances = [r for r in list(workbook['Signal distances'].values)[1:] if r[0]=='NN']
        nn_volumes = [r for r in list(workbook['Cell volumes'].values)[1:] if r[0]=='NN']
        assert len(nn_distances) == len(report['points'])
        assert len(nn_volumes) == len(report['cell_volumes'])
        assert np.isclose(nn_distances[0][-1], report['points'][0]['distance_um'], atol=1e-14, rtol=1e-14)
        automatic_rows = {name: list(workbook[name].values) for name in workbook.sheetnames}
        workbook.close()
        button_path = (output / 'button-export.xlsx').resolve()
        original_dialog = QtWidgets.QFileDialog.getSaveFileName
        QtWidgets.QFileDialog.getSaveFileName = lambda *a, **kw: (str(button_path), 'Excel (*.xlsx)')
        try:
            window.excel_button.click()
            wait(lambda: button_path.exists() and not window.loading)
        finally:
            QtWidgets.QFileDialog.getSaveFileName = original_dialog
        workbook = load_workbook(button_path, data_only=True)
        assert workbook.sheetnames == ['Cell volumes', 'Signal distances']
        assert {name: list(workbook[name].values) for name in workbook.sheetnames} == automatic_rows
        workbook.close()
        assert digest(window.dataset.metadata.path) == original_hash == digest(args.image)
        surfaces = {int(s.name.split()[2]): SurfaceLocator(s) for s in window.dataset.scene.surfaces}
        for point in report['points']:
            locator = surfaces[point['cell_id']]
            b, shortest, _ = locator.closest(point['xyz'])
            assert abs(shortest-point['distance_um']) < 1e-5
            assert locator.closest(point['border_xyz'])[1] < 1e-5
        assert all(m.labels == ['A', 'B'] and m.edges == [(0, 1)] for m in window.dataset.scene.measurements)
        window.tabs.setCurrentIndex(1)
        app.processEvents()
        window.grab().save(str(output / 'distances.png'))
        window.distances.selectRow(0)
        window.tabs.setCurrentIndex(0)
        app.processEvents()
        window.grab().save(str(output / 'result-2d.png'))
        window.mode.setCurrentIndex(1)
        app.processEvents()
        window.focus_measurement()
        app.processEvents()
        assert window.volume_view._annotation_actors
        window.volume_view.plotter.screenshot(str(output / 'result-3d.png'))
        actors = len(window.volume_view._annotation_actors)
        window.show_lines.setChecked(False)
        assert len(window.volume_view._annotation_actors) < actors
        window.show_lines.setChecked(True)
        summary = dict(status='passed', analysis_report=str(window.result_path),
                       package_root=str(args.package_root.resolve()) if args.package_root else None,
                       device=args.device,
                       boundary_model=window.boundary_model.text() or None,
                       signal_model=window.signal_model.text(),
                       membrane_snap=args.membrane_snap,
                       cells=len(report['signal_filter']['kept_cell_ids']), signals=len(report['points']),
                       excel=str(workbook_path), volumes_um3=[r['volume_um3'] for r in report['cell_volumes']],
                       distance_range_um=[min(p['distance_um'] for p in report['points']), max(p['distance_um'] for p in report['points'])],
                       source_sha256=original_hash, native_imaris_runtime_tested=False)
        (output / 'summary.json').write_text(json.dumps(summary, indent=2))
        print(json.dumps(summary, indent=2), flush=True)
        window.close()
        app.exit(0)
    except Exception:
        import traceback
        traceback.print_exc()
        window.close()
        app.exit(1)


QtCore.QTimer.singleShot(100, exercise)
raise SystemExit(app.exec())
