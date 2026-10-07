"""Compare the same detected cells/manual A points before and after refinement."""
import argparse
import csv
import json
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import numpy as np
from openpyxl import Workbook

from fish_viewer.neural.audit import read_references
from fish_viewer.pipeline.evaluation import rasterize_reference
from fish_viewer.pipeline.runner import read_prediction, _json
from fish_viewer.pipeline.segmentation import Prepared
from fish_viewer.signals.data import at_labels,marker_points,reference_memberships
from fish_viewer.workflow.distances import SurfaceLocator
from fish_viewer.workflow.volumes import surface_volume
from benchmark_measurements import manual_meshes,ref_key,paired_stats


def surface_samples(surface):
    triangles=surface.vertices[surface.triangles]
    area=np.linalg.norm(np.cross(triangles[:,1]-triangles[:,0],triangles[:,2]-triangles[:,0]),axis=1)
    valid=area>0;triangles=triangles[valid];area=area[valid]
    rng=np.random.default_rng(42)
    selected=rng.choice(len(triangles),size=256,p=area/area.sum())
    uv=rng.random((256,2));uv[uv.sum(axis=1)>1]=1-uv[uv.sum(axis=1)>1]
    t=triangles[selected]
    return t[:,0]+uv[:,0,None]*(t[:,1]-t[:,0])+uv[:,1,None]*(t[:,2]-t[:,0])


def boundary_error(manual,predicted,manual_locator,predicted_locator,manual_samples):
    distances=np.array([manual_locator.closest(p)[1] for p in surface_samples(predicted)]
                       +[predicted_locator.closest(p)[1] for p in manual_samples])
    return float(distances.mean()),float(np.percentile(distances,95))


def compare(audit,baseline,refined):
    source=json.loads(Path(audit).read_text());old=json.loads(Path(baseline).read_text());new=json.loads(Path(refined).read_text())
    if old['status']!='complete' or new['status']!='complete':raise ValueError('Completed inference is required.')
    old_rows={r['image']:r for r in old['files']};new_rows={r['image']:r for r in new['files']}
    if set(old_rows)!=set(new_rows):raise ValueError('Baseline and refined scans must be identical.')
    sources={s['image']:s for s in source['sources']}
    cells=[];signals=[]
    for number,case in enumerate([c for c in source['cases'] if c['image'] in old_rows],1):
        print(f"[{number}/{len(old_rows)}] Compare boundaries: {Path(case['image']).name}",flush=True)
        old_row=old_rows[case['image']];new_row=new_rows[case['image']]
        old_meshes,metadata=read_prediction(old_row['predictions'])
        new_meshes,new_metadata=read_prediction(new_row['predictions'])
        old_meshes={int(s.name.rsplit(' ',1)[-1]):s for s in old_meshes}
        new_meshes={int(s.name.rsplit(' ',1)[-1]):s for s in new_meshes}
        with np.load(old_row['predictions'],allow_pickle=False) as a:old_labels=a['labels']
        with np.load(new_row['predictions'],allow_pickle=False) as a:new_labels=a['labels']
        if set(old_meshes)!=set(new_meshes):raise ValueError('Refinement changed detected instance identities.')
        references,names=read_references(Path(case['directory'])/'references.npz')
        scenes=sorted({s for alias in case['audit']['aliases'] for s in sources[alias]['scenes']})
        prepared=Prepared(np.empty(old_labels.shape,np.uint8),tuple(metadata['spacing_zyx']),
                          tuple(metadata['origin_xyz']),metadata['channel'],[],metadata['fingerprint'],{},{})
        manual,ab=manual_meshes(scenes,prepared)
        fine=Prepared(np.empty(new_labels.shape,np.uint8),tuple(new_metadata['spacing_zyx']),
                     tuple(new_metadata['origin_xyz']),new_metadata['channel'],[],metadata['fingerprint'],{},{})
        markers=marker_points(scenes)
        membership=reference_memberships([p['xyz'] for p in markers],references,prepared.origin,prepared.spacing)
        sizes=np.bincount(new_labels.ravel())
        diagnostics={r['cell_id']:r for r in new_metadata['refinement']['cells']}
        for ref_id,match in enumerate(old_row['metrics']['per_cell']):
            cell=match['predicted_label']
            if cell is None or match['iou']<.5 or case['audit']['reference_cut'][ref_id]:continue
            truth=manual.get(ref_key(references[ref_id]))
            if truth is None or surface_volume(truth) is None:continue
            refs,_,_=rasterize_reference([truth],fine)
            if len(refs)!=1:continue
            box,mask=refs[0];intersection=int(np.sum(new_labels[box][mask]==cell))
            iou=intersection/(int(mask.sum())+int(sizes[cell])-intersection)
            truth_locator,old_locator,new_locator=SurfaceLocator(truth),SurfaceLocator(old_meshes[cell]),SurfaceLocator(new_meshes[cell])
            samples=surface_samples(truth)
            old_assd,old_hd95=boundary_error(truth,old_meshes[cell],truth_locator,old_locator,samples)
            new_assd,new_hd95=boundary_error(truth,new_meshes[cell],truth_locator,new_locator,samples)
            common=dict(image=str(Path(case['image']).relative_to(ROOT)),group=case['audit']['group'],
                        split=case['audit']['split'],manual_cell=names[ref_id],reference_id=ref_id+1,predicted_cell=cell,
                        baseline_iou=match['iou'],refined_iou=float(iou),refinement_accepted=diagnostics[cell]['accepted'])
            cells.append(dict(common,manual_volume_um3=surface_volume(truth),baseline_volume_um3=surface_volume(old_meshes[cell]),
                              refined_volume_um3=surface_volume(new_meshes[cell]),baseline_surface_mean_error_um=old_assd,
                              refined_surface_mean_error_um=new_assd,baseline_surface_hd95_um=old_hd95,refined_surface_hd95_um=new_hd95))
            for point,members in zip(markers,membership):
                if members!=[ref_id]:continue
                expected=ab.get(tuple(np.round(point['xyz'],4)))
                if expected is None:continue
                old_inside=int(at_labels(old_labels,[point['xyz']],prepared.origin,prepared.spacing)[0])==cell
                new_inside=int(at_labels(new_labels,[point['xyz']],fine.origin,fine.spacing)[0])==cell
                signals.append(dict(common,manual_a_xyz=point['xyz'],manual_ab_distance_um=expected,
                                    baseline_distance_um=old_locator.closest(point['xyz'])[1],
                                    refined_distance_um=new_locator.closest(point['xyz'])[1],
                                    baseline_a_inside=old_inside,refined_a_inside=new_inside))
    eligible=[r for r in signals if r['baseline_a_inside'] and r['refined_a_inside'] and r['refined_iou']>=.5]
    # Keep the original distance cohort visible even when refinement moves a
    # boundary across A. Such crossings are failed eligibility, not outliers.
    original_inside=[r for r in signals if r['baseline_a_inside']]
    summary=dict(scans=len(old_rows),matched_baseline_complete_cells=len(cells),
                 cells_remaining_iou_ge_05=sum(r['refined_iou']>=.5 for r in cells),
                 total_refined_cells=sum(d['accepted'] for r in new['files'] for d in r['refinement']),
                 total_detector_cells=sum(len(r['refinement']) for r in new['files']),
                 eligible_manual_distances=len(eligible),manual_distance_candidates=len(signals),
                 baseline_inside_but_refined_outside=sum(r['baseline_a_inside'] and not r['refined_a_inside'] for r in signals),
                 original_inside_manual_distances=len(original_inside),
                 refined_outside_rate_percent=float(np.mean([not r['refined_a_inside'] for r in original_inside])*100) if original_inside else None,
                 baseline_volume=paired_stats(cells,'manual_volume_um3','baseline_volume_um3'),
                 refined_volume=paired_stats(cells,'manual_volume_um3','refined_volume_um3'),
                 baseline_distance=paired_stats(eligible,'manual_ab_distance_um','baseline_distance_um'),
                 refined_distance=paired_stats(eligible,'manual_ab_distance_um','refined_distance_um'),
                 baseline_surface_mean_error_um=float(np.mean([r['baseline_surface_mean_error_um'] for r in cells])) if cells else None,
                 refined_surface_mean_error_um=float(np.mean([r['refined_surface_mean_error_um'] for r in cells])) if cells else None,
                 baseline_surface_hd95_um=float(np.mean([r['baseline_surface_hd95_um'] for r in cells])) if cells else None,
                 refined_surface_hd95_um=float(np.mean([r['refined_surface_hd95_um'] for r in cells])) if cells else None)
    return dict(status='complete',summary=summary,cells=cells,distances=signals,
                baseline_report=str(Path(baseline).resolve()),refined_report=str(Path(refined).resolve()),
                policy='Fixed baseline complete matches at IoU >= .5, same instance identities and manual A centers. '
                       'Volume/surface errors include all original matches, including degraded matches. '
                       'Primary distance comparison requires A inside both predictions and refined IoU >= .5; '
                       'lost points and all candidate distances remain listed. All outliers included. '
                       'Surface errors sample 256 area-uniform points per surface, seed 42. '
                       'No signal-network localization or signal-retention selection in this boundary comparison.')


def fixed_measurements(benchmark, refined_report, boundary_report):
    """Freeze the original retained cells and CNN signal centers, expose lost eligibility."""
    old=json.loads(Path(benchmark).read_text());new=json.loads(Path(refined_report).read_text())
    meshes={};labels={};metadata={}
    for row in new['files']:
        key=str(Path(row['image']).relative_to(ROOT))
        surfaces,info=read_prediction(row['predictions'])
        meshes[key]={int(s.name.rsplit(' ',1)[-1]):s for s in surfaces}
        metadata[key]=info
        with np.load(row['predictions'],allow_pickle=False) as a:labels[key]=a['labels']
    ious={(r['image'],r['predicted_cell']):r['refined_iou'] for r in boundary_report['cells']}
    cells=[];signals=[]
    for r in old['volumes']:
        if r['iou']<.5 or not r['complete'] or not r['retained_after_signal_filter'] or r['manual_volume_um3'] is None:continue
        key=r['image'];cell=r['predicted_cell']
        cells.append(dict(r,baseline_volume_um3=r['predicted_volume_um3'],
                          refined_volume_um3=surface_volume(meshes[key][cell]),refined_iou=ious[key,cell]))
    for r in old['end_to_end_distances']:
        key=r['image'];cell=r['predicted_cell'];info=metadata[key]
        inside=int(at_labels(labels[key],[r['predicted_a_xyz']],info['origin_xyz'],info['spacing_zyx'])[0])==cell
        locator=SurfaceLocator(meshes[key][cell])
        signals.append(dict(r,baseline_distance_um=r['predicted_distance_um'],
                            refined_distance_um=locator.closest(r['predicted_a_xyz'])[1],
                            predicted_a_inside_refined=inside,refined_iou=ious[key,cell]))
    eligible=[r for r in signals if r['predicted_a_inside_refined'] and r['refined_iou']>=.5]
    summary=dict(original_cells=len(cells),original_signals=len(signals),eligible_signals=len(eligible),
                 signals_outside_refined_cell=sum(not r['predicted_a_inside_refined'] for r in signals),
                 cells_below_iou_05=sum(r['refined_iou']<.5 for r in cells),
                 baseline_volume=paired_stats(cells,'manual_volume_um3','baseline_volume_um3'),
                 refined_volume=paired_stats(cells,'manual_volume_um3','refined_volume_um3'),
                 baseline_distance_same_eligible_signals=paired_stats(eligible,'manual_ab_distance_um','baseline_distance_um'),
                 refined_distance=paired_stats(eligible,'manual_ab_distance_um','refined_distance_um'))
    return dict(summary=summary,cells=cells,signals=signals,policy='Original signal centers and retained matched cells '
                'are frozen; no new signal detection. Volume comparison includes all original 79 cells. '
                'Distances require the original predicted A center to remain inside its assigned cell and '
                'refined IoU >= .5. Lost eligibility is reported; all original 73 signal rows remain available.')


def write_outputs(output,report):
    output.mkdir(parents=True,exist_ok=True);_json(output/'comparison.json',report)
    workbook=Workbook();workbook.remove(workbook.active)
    for name in ('cells','distances'):
        rows=report[name];fields=list(rows[0]) if rows else ['No observations']
        sheet=workbook.create_sheet(name);sheet.append(fields)
        with (output/(name+'.csv')).open('w',newline='',encoding='utf-8') as stream:
            writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
            for row in rows:
                formatted={k:json.dumps(v) if isinstance(v,list) else v for k,v in row.items()}
                writer.writerow(formatted);sheet.append([formatted.get(k) for k in fields])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    sheet=workbook.create_sheet('Summary');sheet.append(['Metric','Value'])
    for k,v in report['summary'].items():sheet.append([k,json.dumps(v)])
    sheet.append(['Policy',report['policy']]);workbook.save(output/'boundary-comparison.xlsx');workbook.close()
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    figure=Figure(figsize=(11,4.5),layout='constrained');FigureCanvasAgg(figure);axes=figure.subplots(1,2)
    cells=report['cells'];distances=[r for r in report['distances'] if r['baseline_a_inside'] and r['refined_a_inside'] and r['refined_iou']>=.5]
    for ax,rows,truth,before,after,title in [
        (axes[0],cells,'manual_volume_um3','baseline_volume_um3','refined_volume_um3','Absolute volume error (%)'),
        (axes[1],distances,'manual_ab_distance_um','baseline_distance_um','refined_distance_um','Absolute distance error (µm)')]:
        relative=truth=='manual_volume_um3'
        x=np.array([abs(r[before]-r[truth])/(r[truth] if relative else 1)*(100 if relative else 1) for r in rows])
        y=np.array([abs(r[after]-r[truth])/(r[truth] if relative else 1)*(100 if relative else 1) for r in rows])
        ax.scatter(x,y,s=18,alpha=.65);limit=max(x.max(initial=1e-6),y.max(initial=1e-6))*1.05
        ax.plot([0,limit],[0,limit],'--',color='gray');ax.set(xlabel='Original detector',ylabel='With refinement',title=f'{title} · {len(rows)} pairs');ax.grid(alpha=.2)
    figure.savefig(output/'boundary-comparison.png',dpi=170)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--audit',type=Path,required=True);parser.add_argument('--baseline-report',type=Path,required=True)
    parser.add_argument('--refined-report',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--baseline-measurements',type=Path,help='Freeze cells and CNN centers from an existing measurement benchmark')
    args=parser.parse_args();report=compare(args.audit,args.baseline_report,args.refined_report)
    if args.baseline_measurements:
        report['fixed_measurements']=fixed_measurements(args.baseline_measurements,args.refined_report,report)
        workbook=Workbook();workbook.remove(workbook.active)
        args.output.mkdir(parents=True,exist_ok=True)
        for name in ('cells','signals'):
            rows=report['fixed_measurements'][name];fields=list(rows[0]) if rows else ['No observations']
            page=workbook.create_sheet(name);page.append(fields)
            with (args.output/('fixed-'+name+'.csv')).open('w',newline='',encoding='utf-8') as stream:
                writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
                for row in rows:
                    values={k:json.dumps(v) if isinstance(v,list) else v for k,v in row.items()}
                    writer.writerow(values);page.append([values.get(k) for k in fields])
        page=workbook.create_sheet('Summary');page.append(['Metric','Value'])
        for key,value in report['fixed_measurements']['summary'].items():page.append([key,json.dumps(value)])
        workbook.save(args.output/'fixed-measurements.xlsx');workbook.close()
    write_outputs(args.output,report);print(json.dumps(report['summary'],indent=2))


if __name__=='__main__':main()
