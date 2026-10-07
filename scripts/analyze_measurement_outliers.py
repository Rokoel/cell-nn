"""Recalculate matched-pair errors with transparent, auditable exclusions."""
import argparse
import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[1]


def statistics(rows, truth_key, prediction_key, tolerances):
    truth = np.array([float(r[truth_key]) for r in rows])
    predicted = np.array([float(r[prediction_key]) for r in rows])
    error = predicted - truth
    absolute = np.abs(error)
    percentage = absolute / truth * 100
    groups = sorted({r['group'] for r in rows})
    totals = np.array([[absolute[[r['group'] == g for r in rows]].sum(),
                        sum(r['group'] == g for r in rows)] for g in groups])
    indices = np.random.default_rng(42).integers(len(groups), size=(2000, len(groups)))
    draws = totals[indices].sum(axis=1)
    ci = np.percentile(draws[:, 0] / draws[:, 1], [2.5, 97.5]).tolist()
    return dict(n=len(rows), scans=len({r['image'] for r in rows}), groups=len(groups),
                mae=float(absolute.mean()), mse=float(np.mean(error**2)),
                rmse=float(np.sqrt(np.mean(error**2))), bias=float(error.mean()),
                median_absolute_error=float(np.median(absolute)),
                p95_absolute_error=float(np.percentile(absolute, 95)),
                mean_absolute_percentage_error=float(percentage.mean()),
                median_absolute_percentage_error=float(np.median(percentage)),
                mean_signed_percentage_error=float(np.mean(error / truth * 100)),
                mae_95ci_acquisition_bootstrap=ci,
                within_tolerances={str(t): dict(count=int(np.sum(values <= t)),
                                               percent=float(np.mean(values <= t)*100))
                                   for t, values in tolerances(absolute, percentage)})


def analyze(rows, truth_key, prediction_key, relative):
    truth = np.array([float(r[truth_key]) for r in rows])
    predicted = np.array([float(r[prediction_key]) for r in rows])
    if not np.isfinite(truth).all() or not np.isfinite(predicted).all() or np.any(truth <= 0):
        raise ValueError('All pairs must be finite with positive reference measurements.')
    absolute = np.abs(predicted - truth)
    percentage = absolute / truth * 100
    severity = percentage if relative else absolute
    q1, q3 = np.percentile(severity, [25, 75])
    fence = q3 + 1.5*(q3-q1)
    iqr_keep = severity <= fence
    trim_count = int(np.ceil(len(rows)*.05))
    trim_keep = np.ones(len(rows), dtype=bool)
    trim_keep[np.argsort(severity, kind='stable')[-trim_count:]] = False
    tolerances = (lambda a, p: [(t, p) for t in (10, 20)]) if relative else (
        lambda a, p: [(t, a) for t in (.1, .2, .3)])
    result = dict(error_for_exclusion='absolute_percentage_error' if relative else 'absolute_error_um',
                  q1=float(q1), q3=float(q3), iqr_upper_fence=float(fence),
                  iqr_excluded=int(np.sum(~iqr_keep)), trim_5_percent_excluded=trim_count,
                  tolerance_unit='percent' if relative else 'um', scenarios={})
    for name, mask in [('all', np.ones(len(rows), dtype=bool)),
                       ('iqr_outliers_excluded', iqr_keep), ('worst_5_percent_trimmed', trim_keep)]:
        result['scenarios'][name] = statistics([r for r, keep in zip(rows, mask) if keep],
                                              truth_key, prediction_key, tolerances)
    audit = [dict(**r, signed_error=float(p-m), absolute_error=float(a),
                  absolute_percentage_error=float(pct), iqr_outlier=bool(not ik),
                  excluded_by_5_percent_trim=bool(not tk))
             for r, m, p, a, pct, ik, tk in zip(rows, truth, predicted, absolute, percentage,
                                               iqr_keep, trim_keep)]
    return result, audit


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=ROOT/'reports/measurement-accuracy')
    parser.add_argument('--output', type=Path, default=ROOT/'reports/measurement-accuracy/outlier-analysis')
    args = parser.parse_args()
    paths = [args.input/'volumes.csv', args.input/'end_to_end_distances.csv']
    tables = []
    for path in paths:
        with path.open(newline='', encoding='utf-8') as stream:
            tables.append(list(csv.DictReader(stream)))
    cells = [r for r in tables[0] if float(r['iou']) >= .5 and r['complete'].lower() == 'true'
             and r['retained_after_signal_filter'].lower() == 'true'
             and r['manual_volume_um3'] and r['predicted_volume_um3']]
    signals = [r for r in tables[1] if r['manual_ab_distance_um'] and r['predicted_distance_um']]
    volume, cell_audit = analyze(cells, 'manual_volume_um3', 'predicted_volume_um3', True)
    distance, signal_audit = analyze(signals, 'manual_ab_distance_um', 'predicted_distance_um', False)
    report = dict(method='Upper Tukey fence: Q3 + 1.5 * (Q3 - Q1), NumPy linear percentiles.',
                  secondary_method='Remove ceil(0.05*n) largest errors; fixed-percentage sensitivity only.',
                  source_sha256={p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},
                  bootstrap='2000 acquisition-group resamples, seed 42; exclusions held fixed.',
                  caveat='Conditional on matched detections and retained observations; reused test split. '
                         'Exclusions do not correct failed measurements or establish independent accuracy.',
                  volume=volume, distance=distance)
    args.output.mkdir(parents=True, exist_ok=True)
    (args.output/'summary.json').write_text(json.dumps(report, indent=2)+'\n', encoding='utf-8')
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append(['Measurement', 'Scenario', 'N', 'MAE', 'RMSE', 'Bias', 'Median absolute error',
                    'Mean absolute percentage error', 'Median absolute percentage error',
                    'P95 absolute error', 'MAE CI lower', 'MAE CI upper'])
    for name, result in [('Volume (um3)', volume), ('Distance (um)', distance)]:
        for scenario, s in result['scenarios'].items():
            summary.append([name, scenario, s['n'], s['mae'], s['rmse'], s['bias'],
                            s['median_absolute_error'], s['mean_absolute_percentage_error'],
                            s['median_absolute_percentage_error'], s['p95_absolute_error'],
                            *s['mae_95ci_acquisition_bootstrap']])
    for name, audit in [('Volumes', cell_audit), ('Distances', signal_audit)]:
        fields = list(audit[0])
        sheet = book.create_sheet(name)
        sheet.append(fields)
        for row in audit:
            sheet.append([row[f] for f in fields])
        with (args.output/(name.lower()+'-audit.csv')).open('w', newline='', encoding='utf-8') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(audit)
    methods = book.create_sheet('Methods')
    for key in ('method', 'secondary_method', 'bootstrap', 'caveat'):
        methods.append([key, report[key]])
    methods.append(['Volume IQR fence (%)', volume['iqr_upper_fence']])
    methods.append(['Distance IQR fence (um)', distance['iqr_upper_fence']])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(14,
                max(len(str(c.value or '')) for c in column)+2))
    book.save(args.output/'outlier-analysis.xlsx')
    lines = ['# Matched-measurement outlier sensitivity', '', report['method'], '',
             'Volume exclusions use absolute percentage error; distance exclusions use absolute error in µm.',
             f"Volume upper fence: {volume['iqr_upper_fence']:.3f}%; excluded {volume['iqr_excluded']}/{len(cells)}.",
             f"Distance upper fence: {distance['iqr_upper_fence']:.6f} µm; excluded {distance['iqr_excluded']}/{len(signals)}.", '',
             '| Measurement | Scenario | N | MAE | RMSE | Mean absolute % error | Median absolute % error |',
             '|---|---|---:|---:|---:|---:|---:|']
    for name, result in [('Volume (µm³)', volume), ('Distance (µm)', distance)]:
        for scenario, s in result['scenarios'].items():
            lines.append(f"| {name} | {scenario.replace('_', ' ')} | {s['n']} | {s['mae']:.4f} | "
                         f"{s['rmse']:.4f} | {s['mean_absolute_percentage_error']:.2f}% | "
                         f"{s['median_absolute_percentage_error']:.2f}% |")
    lines.extend(['', '## Tolerance coverage', ''])
    for name, result in [('Volume', volume), ('Distance', distance)]:
        for scenario, s in result['scenarios'].items():
            coverage = '; '.join(f"within {t} {result['tolerance_unit']}: {v['count']}/{s['n']} "
                                 f"({v['percent']:.1f}%)" for t, v in s['within_tolerances'].items())
            lines.append(f"- {name}, {scenario.replace('_', ' ')}: {coverage}.")
    lines.extend(['', '## Excluded by the standard IQR rule', ''])
    for name, audit in [('Volume', cell_audit), ('Distance', signal_audit)]:
        excluded = [r for r in audit if r['iqr_outlier']]
        if not excluded:
            lines.append(f'- {name}: none.')
        for row in excluded:
            lines.append(f"- {name}: `{row['image']}`, predicted cell {row['predicted_cell']}, "
                         f"signal {row.get('signal_id', 'n/a')}; absolute error {row['absolute_error']:.6f}.")
    lines.extend(['', 'The 5% sensitivity removes four cells and four signals separately, ranked using the same',
                  'error definitions as the IQR rule. It does not mean all four are statistical outliers.',
                  'The cell and signal comparisons remain separate cohorts.', '',
                  'Percentage error = 100 × |predicted − manual| / manual. Distance percentages are unstable',
                  'for very short manual distances; interpret distance MAE in µm first.', '',
                  report['bootstrap'], 'These intervals condition on the selected observations and do not',
                  'include uncertainty in outlier selection or manual annotation.', '', report['caveat'], '',
                  'Original CSVs, unfiltered results and graphs are preserved. Audit CSVs and workbook',
                  'contain every original primary pair with explicit exclusion flags.', '',
                  'Reproduce: `.venv/bin/python scripts/analyze_measurement_outliers.py`', ''])
    (args.output/'README.md').write_text('\n'.join(lines), encoding='utf-8')
    print(json.dumps({'volume': volume, 'distance': distance}, indent=2))


if __name__ == '__main__':
    main()
