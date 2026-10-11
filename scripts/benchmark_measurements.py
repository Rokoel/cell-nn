"""Evaluate fixed cell/signal checkpoints on the original held-out acquisition split.

Run cell inference first with --evidence-only (or pass its running report with
--wait). This script scores the signal CNN once loaded, without using manual
annotations for inference. Reference annotations are used only for comparison.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import numpy as np
from scipy.optimize import linear_sum_assignment
from scipy.spatial.distance import cdist
from openpyxl import Workbook

from fish_viewer.worker_runtime import restart_with_cuda
from fish_viewer.models import LibraryEntry
from fish_viewer.reader import load_dataset
from fish_viewer.scenes import read_scene
from fish_viewer.neural.audit import read_references
from fish_viewer.pipeline.evaluation import evaluate, rasterize_reference
from fish_viewer.pipeline.segmentation import Prepared
from fish_viewer.pipeline.runner import read_prediction, _json
from fish_viewer.signals.config import SignalConfig
from fish_viewer.signals.data import normalized, native_labels, candidates, extract_patches, marker_points, infer_channels, reference_memberships, at_labels
from fish_viewer.signals.engine import network, predict_patches, image_fingerprint
from fish_viewer.signals.evaluation import suppress
from fish_viewer.neural.engine import runtime
from fish_viewer.workflow.distances import SurfaceLocator
from fish_viewer.workflow.volumes import surface_volume


def ref_key(reference):
    box, mask = reference
    return hashlib.sha256(str([(s.start,s.stop) for s in box]).encode()+mask.tobytes()).hexdigest()


def paired_stats(rows, expected, predicted):
    rows = [r for r in rows if r.get(expected) is not None and r.get(predicted) is not None]
    if not rows:
        return dict(n=0)
    truth = np.array([r[expected] for r in rows])
    found = np.array([r[predicted] for r in rows])
    error = found-truth
    result = dict(n=len(rows), bias=float(error.mean()), mae=float(np.abs(error).mean()),
                  mse=float(np.mean(error**2)), rmse=float(np.sqrt(np.mean(error**2))),
                  median_absolute_error=float(np.median(np.abs(error))),
                  p95_absolute_error=float(np.percentile(np.abs(error),95)))
    if np.all(truth > 0):
        result['mean_absolute_percentage_error'] = float(np.mean(np.abs(error)/truth)*100)
        result['median_absolute_percentage_error'] = float(np.median(np.abs(error)/truth)*100)
    groups = sorted({r['group'] for r in rows})
    if len(groups) > 1:
        sums = np.array([[sum(abs(r[predicted]-r[expected]) for r in rows if r['group']==g),
                          sum(r['group']==g for r in rows)] for g in groups])
        indices = np.random.default_rng(42).integers(len(groups), size=(2000,len(groups)))
        samples = sums[indices].sum(axis=1)
        result['mae_95ci_acquisition_bootstrap'] = np.percentile(samples[:,0]/samples[:,1], [2.5,97.5]).tolist()
    return result


def manual_meshes(scenes, prepared):
    meshes = {}
    ab = {}
    for path in scenes:
        scene = read_scene(Path(path))
        for surface in scene.surfaces:
            references, _, _ = rasterize_reference([surface], prepared)
            if references:
                meshes.setdefault(ref_key(references[0]), surface)
        for group in scene.measurements:
            for left, right in group.edges:
                if group.times[left] != 0 or group.times[right] != 0:
                    continue
                labels = [group.labels[left].strip().upper(),group.labels[right].strip().upper()]
                if set(labels) != {'A','B'}:
                    continue
                a,b = (left,right) if labels[0]=='A' else (right,left)
                key = tuple(np.round(group.points[a],4))
                ab.setdefault(key,[]).append(float(np.linalg.norm(group.points[a]-group.points[b])))
    # Conflicting annotations on duplicate scans are reported as ambiguous.
    return meshes, {k: float(np.mean(v)) if np.ptp(v)<=.05 else None for k,v in ab.items()}


def analyze(case, row, collection, model, config, threshold, output):
    with np.load(row['predictions'],allow_pickle=False) as archive:
        labels = archive['labels'].copy()
        metadata = json.loads(str(archive['metadata']))
    if metadata['config'].get('use_shape_filter', True):
        raise ValueError('Benchmark requires the current app policy with shape filtering disabled.')
    dataset = load_dataset(LibraryEntry(Path(case['image'])))
    if image_fingerprint(dataset) != metadata['fingerprint']:
        raise ValueError('Source pixels differ from predictions.')
    images = [normalized(v) for v in dataset.volumes]
    spacing = np.asarray(dataset.spacing[::-1]); origin = np.asarray(dataset.center_origin)
    membrane = int(metadata['channel'])
    native = native_labels(labels, metadata, dataset)
    proposed = [dict(p,channel=c) for c,volume in enumerate(images) if c!=membrane
                for p in candidates(volume,native,spacing,origin,config)]
    scores = []
    for start in range(0,len(proposed),64):
        patches = extract_patches(images,membrane,native,proposed[start:start+64],origin,spacing,config)
        scores.extend(predict_patches(model,patches).tolist())
    points = [dict(p,signal_id=i+1) for i,p in enumerate(suppress(proposed,scores,threshold,config.candidate_separation_um))]
    kept = {p['cell_id'] for p in points}
    references,names = read_references(Path(case['directory'])/'references.npz')
    metrics = evaluate(labels,references,names,metadata['spacing_zyx'])
    sources = {s['image']:s for s in collection['sources']}
    scenes = sorted({s for alias in case['audit']['aliases'] for s in sources[alias]['scenes']})
    prepared = Prepared(np.load(Path(case['directory'])/'image.npy',mmap_mode='r'),
                        tuple(metadata['spacing_zyx']),tuple(metadata['origin_xyz']),membrane,[],metadata['fingerprint'],{}, {})
    meshes, ab = manual_meshes(scenes,prepared)
    prediction_meshes, _ = read_prediction(row['predictions'])
    nn_meshes = {int(s.name.rsplit(' ',1)[-1]):s for s in prediction_meshes}
    markers = infer_channels(images,marker_points(scenes),membrane,origin,spacing,config)
    memberships = reference_memberships([p['xyz'] for p in markers],references,metadata['origin_xyz'],metadata['spacing_zyx'])
    for p,members in zip(markers,memberships):
        p['manual_memberships'] = members
    common = dict(image=str(Path(case['image']).relative_to(ROOT)),group=case['audit']['group'])
    volume_rows, contour_rows, end_rows = [],[],[]
    matched_reference_map = {}
    for ref_id, match in enumerate(metrics['per_cell']):
        cell = match['predicted_label']
        if cell is None:
            continue
        manual_mesh = meshes.get(ref_key(references[ref_id]))
        nn_mesh = nn_meshes[cell]
        complete = not case['audit']['reference_cut'][ref_id]
        manual_volume = surface_volume(manual_mesh) if manual_mesh else None
        nn_volume = surface_volume(nn_mesh)
        v = dict(common,manual_cell=names[ref_id],reference_id=ref_id+1,predicted_cell=cell,
                 iou=match['iou'],dice=match['dice'],complete=complete,
                 retained_after_signal_filter=cell in kept,manual_volume_um3=manual_volume,predicted_volume_um3=nn_volume)
        volume_rows.append(v)
        if match['iou'] < .5 or not complete or manual_mesh is None:
            continue
        matched_reference_map[ref_id] = cell
        nn_locator, manual_locator = SurfaceLocator(nn_mesh),SurfaceLocator(manual_mesh)
        for p in markers:
            if p['manual_memberships'] != [ref_id]:
                continue
            _, truth_distance, _ = manual_locator.closest(p['xyz'])
            _, found_distance, _ = nn_locator.closest(p['xyz'])
            inside = int(at_labels(labels,[p['xyz']],metadata['origin_xyz'],metadata['spacing_zyx'])[0]) == cell
            contour_rows.append(dict(common,manual_cell=names[ref_id],predicted_cell=cell,
                channel=p['channel']+1 if p['channel'] is not None else None,
                retained_after_signal_filter=cell in kept,manual_signal_inside_predicted_cell=inside,
                manual_a_xyz=p['xyz'],manual_mesh_distance_um=truth_distance,
                manual_ab_distance_um=ab.get(tuple(np.round(p['xyz'],4))),
                predicted_border_distance_at_manual_a_um=found_distance))
    # Signal localization is matched independently in each cell and inferred channel.
    truth = [dict(p,cell_id=matched_reference_map[p['manual_memberships'][0]]) for p in markers
             if len(p['manual_memberships'])==1 and p['manual_memberships'][0] in matched_reference_map and p['channel'] is not None]
    selected = [p for p in points if p['cell_id'] in matched_reference_map.values()]
    if selected and truth:
        distance = cdist([p['xyz'] for p in selected],[p['xyz'] for p in truth])
        compatible = np.array([[p['cell_id']==r['cell_id'] and p['channel']==r['channel'] for r in truth] for p in selected])
        cost = np.where(compatible & (distance <= config.match_radius_um),distance,1e6)
        left,right = linear_sum_assignment(cost)
        for a,b in zip(left,right):
            if cost[a,b] >= 1e6:
                continue
            predicted,manual = selected[a],truth[b]
            cell = predicted['cell_id']; ref_id = manual['manual_memberships'][0]
            _,manual_distance,_ = SurfaceLocator(meshes[ref_key(references[ref_id])]).closest(manual['xyz'])
            border,found_distance,_ = SurfaceLocator(nn_meshes[cell]).closest(predicted['xyz'])
            end_rows.append(dict(common,manual_cell=names[ref_id],predicted_cell=cell,channel=predicted['channel']+1,
                signal_id=predicted['signal_id'],score=predicted['score'],localization_error_um=float(distance[a,b]),
                manual_a_xyz=manual['xyz'],predicted_a_xyz=predicted['xyz'],predicted_b_xyz=border.tolist(),
                manual_mesh_distance_um=manual_distance,manual_ab_distance_um=ab.get(tuple(np.round(manual['xyz'],4))),
                predicted_distance_um=found_distance))
    record = dict(common,references=len(references),complete_references=sum(not c for c in case['audit']['reference_cut']),
                  predicted_cells=len(nn_meshes),retained_cells=len(kept),predicted_signals=len(points),
                  candidates=len(proposed),manual_markers=len(markers),ambiguous_channel_markers=sum(p['channel'] is None for p in markers),
                  ambiguous_or_outside_manual_cells=sum(len(m)!=1 for m in memberships),
                  matched_complete_cells=len(matched_reference_map),eligible_manual_signals=len(truth),matched_signals=len(end_rows))
    result = dict(scan=record,volumes=volume_rows,contour_distances=contour_rows,end_to_end_distances=end_rows,predicted_signals=points)
    _json(output/(Path(case['image']).stem+'-'+metadata['fingerprint'][:8]+'.json'),result)
    return result


def outputs(directory, report):
    workbook = Workbook(); workbook.remove(workbook.active)
    for key in ('scans','volumes','contour_distances','end_to_end_distances'):
        rows = report[key]
        fields = list(dict.fromkeys(k for row in rows for k in row))
        page = workbook.create_sheet(key[:31]); page.append(fields or ['No observations'])
        with (directory/(key+'.csv')).open('w',newline='',encoding='utf-8') as stream:
            writer = csv.DictWriter(stream,fields);writer.writeheader()
            for row in rows:
                formatted={k:json.dumps(v) if isinstance(v,(list,dict)) else v for k,v in row.items()}
                writer.writerow(formatted);page.append([formatted.get(k) for k in fields])
                for cell in page[page.max_row]:
                    if isinstance(cell.value,str):cell.data_type='s'
        page.freeze_panes='A2';page.auto_filter.ref=page.dimensions
    page = workbook.create_sheet('Summary');page.append(['Metric','Result'])
    for k,v in report['summary'].items():page.append([k,json.dumps(v)])
    workbook.save(directory/'measurement-accuracy.xlsx');workbook.close()
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    fig=Figure(figsize=(12,4),layout='constrained');FigureCanvasAgg(fig)
    for ax,(rows,x,y,title) in zip(fig.subplots(1,3),[
        ([r for r in report['volumes'] if r['iou']>=.5 and r['complete'] and r['retained_after_signal_filter']],'manual_volume_um3','predicted_volume_um3','Cell volume (µm³)'),
        ([r for r in report['contour_distances'] if r['retained_after_signal_filter'] and r['manual_signal_inside_predicted_cell']],'manual_ab_distance_um','predicted_border_distance_at_manual_a_um','Contour only: distance (µm)'),
        (report['end_to_end_distances'],'manual_ab_distance_um','predicted_distance_um','Full pipeline: distance (µm)')]):
        rows=[r for r in rows if r.get(x) is not None and r.get(y) is not None]
        ax.scatter([r[x] for r in rows],[r[y] for r in rows],s=16,alpha=.65)
        limit=max([r[k] for r in rows for k in (x,y)],default=1)*1.05
        ax.plot([0,limit],[0,limit],'--',color='gray');ax.set(xlabel='Manual',ylabel='Predicted',title=title);ax.grid(alpha=.2)
    fig.savefig(directory/'measurement-accuracy.png',dpi=170)


def refresh_volumes(report, collection, cells_run):
    """Re-integrate saved matches after geometry fixes without repeating CNN inference."""
    rows={r['image']:r for r in cells_run['files']}
    for case in collection['cases']:
        if case['audit']['split']!='test':continue
        key=str(Path(case['image']).relative_to(ROOT))
        matched=[r for r in report['volumes'] if r['image']==key]
        if not matched:continue
        surfaces,metadata=read_prediction(rows[case['image']]['predictions'])
        nn={int(s.name.rsplit(' ',1)[-1]):s for s in surfaces}
        prepared=Prepared(np.load(Path(case['directory'])/'image.npy',mmap_mode='r'),tuple(metadata['spacing_zyx']),
                          tuple(metadata['origin_xyz']),metadata['channel'],[],metadata['fingerprint'],{}, {})
        sources={s['image']:s for s in collection['sources']}
        scenes=sorted({s for alias in case['audit']['aliases'] for s in sources[alias]['scenes']})
        meshes,_=manual_meshes(scenes,prepared)
        references,_=read_references(Path(case['directory'])/'references.npz')
        for r in matched:
            manual=meshes.get(ref_key(references[r['reference_id']-1]))
            r['manual_volume_um3']=surface_volume(manual) if manual else None
            r['predicted_volume_um3']=surface_volume(nn[r['predicted_cell']])
    report['volume_integration']='Closed cleaned triangle polygons; consistent orientation; degenerate zero-area lines removed.'


def main():
    restart_with_cuda()
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--audit',type=Path,required=True)
    parser.add_argument('--cells-report',type=Path,required=True)
    parser.add_argument('--signal-model',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--device',choices=['gpu','cpu','auto'],default='auto')
    parser.add_argument('--wait',action='store_true')
    parser.add_argument('--refresh-volumes',action='store_true',help='Reintegrate volumes in an existing benchmark; reuse saved cell/signal matches')
    args=parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    cache=args.output/'scan-results';cache.mkdir(exist_ok=True)
    collection=json.loads(args.audit.read_text())
    cases=[c for c in collection['cases'] if c['audit']['split']=='test']
    if not cases:raise ValueError('No original held-out test scans found.')
    manifest=json.loads((args.signal_model/'signal-model.json').read_text())
    config=SignalConfig(**manifest['config']).validate()
    if not args.refresh_volumes:
        tf,_,_,info=runtime(args.device,config.seed);model=network(tf,config)
        model.load_weights(args.signal_model/manifest['checkpoint'])
    else:
        info={}
    report=dict(version=1,status='running',policy='Fixed checkpoints and thresholds; no tuning on test. One-to-one whole-cell IoU >= 0.5. Primary: complete, closed manual cells retained by the final any-signal filter. Signal matching: same cell/channel, <=0.6 um. Manual signal membership must be unique. Distance comparison uses stored manual A-B lines; manual-mesh shortest distances are also reported.',
        limitations='Exploratory reused test split, inspected during earlier development. Incomplete cell annotations cannot establish whole-field cell precision. Conditional error on matched objects does not measure missed cells/signals. Open/nonmanifold reference meshes excluded from volume summaries. Channel-ambiguous markers excluded from end-to-end matching.',
        test_scans_expected=len(cases),runtime=info,signal_threshold=manifest['threshold'],
        cell_report=str(args.cells_report.resolve()),signal_model=str(args.signal_model.resolve()),
        scans=[],volumes=[],contour_distances=[],end_to_end_distances=[],errors=[])
    done=set()
    if args.refresh_volumes:
        report=json.loads((args.output/'report.json').read_text())
        if report['status']!='complete':raise ValueError('A completed benchmark is required to refresh volumes.')
        refresh_volumes(report,collection,json.loads(args.cells_report.read_text()))
        done={c['image'] for c in cases}
    while len(done)<len(cases):
        cells_run=json.loads(args.cells_report.read_text());rows={r['image']:r for r in cells_run['files']}
        for case in cases:
            if case['image'] in done or case['image'] not in rows:continue
            print(f"[{len(done)+1}/{len(cases)}] Signals and measurements: {Path(case['image']).name}",flush=True)
            try:
                result=analyze(case,rows[case['image']],collection,model,config,manifest['threshold'],cache)
                report['scans'].append(result['scan'])
                for key in ('volumes','contour_distances','end_to_end_distances'):report[key].extend(result[key])
            except Exception as exc:
                import traceback
                report['errors'].append(dict(image=case['image'],error=str(exc),traceback=traceback.format_exc()))
                print(traceback.format_exc(),flush=True)
            done.add(case['image']);_json(args.output/'report.json',report)
        if cells_run['status'] in ('failed','cancelled'):
            raise RuntimeError('Cell inference did not complete: '+cells_run.get('error','cancelled'))
        if len(done)<len(cases):
            if not args.wait or cells_run['status']=='complete':raise ValueError('Missing held-out scan predictions.')
            time.sleep(3)
    volumes=[r for r in report['volumes'] if r['iou']>=.5 and r['complete']]
    final_volumes=[r for r in volumes if r['retained_after_signal_filter']]
    contours=[r for r in report['contour_distances'] if r['retained_after_signal_filter'] and r['manual_signal_inside_predicted_cell']]
    report['summary']=dict(test_scans_processed=len(report['scans']),
        manual_complete_cells=sum(r['complete_references'] for r in report['scans']),
        matched_complete_cells=len(volumes),matched_retained_complete_cells=len(final_volumes),
        predicted_cells=sum(r['predicted_cells'] for r in report['scans']),
        retained_cells=sum(r['retained_cells'] for r in report['scans']),
        predicted_signals=sum(r['predicted_signals'] for r in report['scans']),
        eligible_manual_signals=sum(r['eligible_manual_signals'] for r in report['scans']),
        matched_signals=len(report['end_to_end_distances']),
        volume_final_pipeline=paired_stats(final_volumes,'manual_volume_um3','predicted_volume_um3'),
        volume_before_signal_filter=paired_stats(volumes,'manual_volume_um3','predicted_volume_um3'),
        volume_loose_iou_01_sensitivity=paired_stats([r for r in report['volumes'] if r['complete'] and r['retained_after_signal_filter']],'manual_volume_um3','predicted_volume_um3'),
        contour_distance_vs_manual_ab=paired_stats(contours,'manual_ab_distance_um','predicted_border_distance_at_manual_a_um'),
        contour_distance_vs_manual_mesh=paired_stats(contours,'manual_mesh_distance_um','predicted_border_distance_at_manual_a_um'),
        end_to_end_distance_vs_manual_ab=paired_stats(report['end_to_end_distances'],'manual_ab_distance_um','predicted_distance_um'),
        end_to_end_distance_vs_manual_mesh=paired_stats(report['end_to_end_distances'],'manual_mesh_distance_um','predicted_distance_um'),
        manual_ab_vs_manual_mesh=paired_stats(report['contour_distances'],'manual_ab_distance_um','manual_mesh_distance_um'),
        invalid_manual_volume_meshes=sum(r['manual_volume_um3'] is None for r in final_volumes))
    report['status']='complete' if not report['errors'] else 'completed_with_errors'
    _json(args.output/'report.json',report);outputs(args.output,report)
    print(json.dumps(report['summary'],indent=2),flush=True)
    return 0 if not report['errors'] else 1


if __name__=='__main__':raise SystemExit(main())
