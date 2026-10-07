"""Compare measurement agreement on shared manual objects across different detectors."""
import argparse
import csv
import json
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import numpy as np
from openpyxl import Workbook
from benchmark_measurements import paired_stats
from fish_viewer.pipeline.runner import _json


def compare(baseline,new):
    def volumes(report):
        return {(r['image'],r['manual_cell']):r for r in report['volumes'] if r['complete'] and r['iou']>=.5
                and r['retained_after_signal_filter'] and r['manual_volume_um3'] is not None}
    def signals(report):
        return {(r['image'],r['channel'],tuple(np.round(r['manual_a_xyz'],4))):r for r in report['end_to_end_distances']}
    old_vol,new_vol=volumes(baseline),volumes(new);old_sig,new_sig=signals(baseline),signals(new)
    cells=[];points=[]
    for key in sorted(old_vol.keys()&new_vol.keys()):
        a,b=old_vol[key],new_vol[key]
        if not np.isclose(a['manual_volume_um3'],b['manual_volume_um3'],rtol=1e-5,atol=1e-5):
            raise ValueError('Manual reference differs on shared cell: '+str(key))
        cells.append(dict(image=a['image'],group=a['group'],manual_cell=a['manual_cell'],
            manual_volume_um3=a['manual_volume_um3'],baseline_volume_um3=a['predicted_volume_um3'],
            new_volume_um3=b['predicted_volume_um3'],baseline_iou=a['iou'],new_iou=b['iou']))
    for key in sorted(old_sig.keys()&new_sig.keys()):
        a,b=old_sig[key],new_sig[key]
        if not np.isclose(a['manual_ab_distance_um'],b['manual_ab_distance_um'],rtol=1e-5,atol=1e-5):
            raise ValueError('Manual A/B reference differs: '+str(key))
        points.append(dict(image=a['image'],group=a['group'],channel=a['channel'],manual_a_xyz=a['manual_a_xyz'],
            manual_ab_distance_um=a['manual_ab_distance_um'],baseline_distance_um=a['predicted_distance_um'],
            new_distance_um=b['predicted_distance_um']))
    summary=dict(baseline_matched_retained_cells=len(old_vol),new_matched_retained_cells=len(new_vol),
        shared_cells=len(cells),baseline_only_cells=len(old_vol.keys()-new_vol.keys()),new_only_cells=len(new_vol.keys()-old_vol.keys()),
        baseline_matched_signals=len(old_sig),new_matched_signals=len(new_sig),shared_signals=len(points),
        baseline_only_signals=len(old_sig.keys()-new_sig.keys()),new_only_signals=len(new_sig.keys()-old_sig.keys()),
        baseline_volume=paired_stats(cells,'manual_volume_um3','baseline_volume_um3'),
        new_volume=paired_stats(cells,'manual_volume_um3','new_volume_um3'),
        baseline_distance=paired_stats(points,'manual_ab_distance_um','baseline_distance_um'),
        new_distance=paired_stats(points,'manual_ab_distance_um','new_distance_um'))
    return dict(status='complete',summary=summary,cells=cells,signals=points,
        policy='Intersection of retained complete manual-cell matches and matched manual signal identities. '
        'Cell matches require whole-cell IoU>=.5 in both pipelines. No outlier removal. '
        'Baseline-only and new-only matches are counted, not interpreted as background. '
        'This reused historical test split is exploratory and cannot measure precision with sparse labels.')


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--baseline',type=Path,required=True)
    p.add_argument('--new',type=Path,required=True);p.add_argument('--output',type=Path,required=True);args=p.parse_args()
    baseline=json.loads(args.baseline.read_text());new=json.loads(args.new.read_text())
    if baseline['status']!='complete' or new['status']!='complete':raise ValueError('Completed benchmarks required.')
    report=compare(baseline,new);args.output.mkdir(parents=True,exist_ok=True);_json(args.output/'comparison.json',report)
    workbook=Workbook();workbook.remove(workbook.active)
    for name in ('cells','signals'):
        rows=report[name];fields=list(rows[0]) if rows else ['No observations'];sheet=workbook.create_sheet(name);sheet.append(fields)
        with (args.output/(name+'.csv')).open('w',newline='',encoding='utf-8') as stream:
            writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
            for r in rows:
                row={k:json.dumps(v) if isinstance(v,list) else v for k,v in r.items()};writer.writerow(row)
                sheet.append([row.get(k) for k in fields])
    sheet=workbook.create_sheet('Summary');sheet.append(['Metric','Value'])
    for k,v in report['summary'].items():sheet.append([k,json.dumps(v)])
    sheet.append(['Policy',report['policy']]);workbook.save(args.output/'comparison.xlsx');workbook.close()
    print(json.dumps(report['summary'],indent=2))


if __name__=='__main__':main()
